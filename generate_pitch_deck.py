import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    # Initialize Presentation in 16:9 Widescreen
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette Definitions
    C_NAVY_DARK = RGBColor(11, 19, 43)        # #0B132B
    C_NAVY_CARD = RGBColor(22, 33, 62)        # #16213E
    C_NAVY_BORDER = RGBColor(42, 60, 95)      # #2A3C5F
    C_INDIGO_ACCENT = RGBColor(79, 70, 229)   # #4F46E5
    C_INDIGO_LIGHT = RGBColor(99, 102, 241)   # #6366F1
    C_CYAN_ACCENT = RGBColor(6, 182, 212)     # #06B6D4
    C_RED_ACCENT = RGBColor(220, 38, 38)      # #DC2626
    C_GREEN_ACCENT = RGBColor(16, 185, 129)   # #10B981
    C_BG_LIGHT = RGBColor(248, 250, 252)      # #F8FAFC
    C_CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF
    C_CARD_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
    C_TEXT_DARK = RGBColor(15, 23, 42)        # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
    C_TEXT_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    C_TEXT_LIGHT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    C_PILL_BG = RGBColor(238, 242, 255)       # #EEF2FF

    IMG_DIR = "/Users/mahak/.gemini/antigravity/brain/63ff7a6c-88ed-420f-957e-9fb7f22791b3/.user_uploaded"
    IMG_LINKED_PLATFORMS = os.path.join(IMG_DIR, "media_1792025167579_8b5562d1.jpg")
    IMG_CASE_DETAIL = os.path.join(IMG_DIR, "media_1792025173258_a059dd31.jpg")
    IMG_PDF_REPORT = os.path.join(IMG_DIR, "media_1792025180232_bf60218e.jpg")
    IMG_SELLER_NETWORK = os.path.join(IMG_DIR, "media_1792025186557_c3749f01.jpg")
    IMG_RAJASTHAN_MAP = os.path.join(IMG_DIR, "media_1792001751224_82fa25ec.png")
    IMG_DASHBOARD = os.path.join(IMG_DIR, "media_1791983658833_1bf808e0.png")
    IMG_CHECK_POST = os.path.join(IMG_DIR, "media_1791987346997_a18630e3.png")

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, tag_text, title_text, subtitle_text="", is_dark=False):
        # Category Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.8), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = C_NAVY_CARD if is_dark else C_PILL_BG
        pill.line.color.rgb = C_INDIGO_LIGHT if is_dark else C_INDIGO_ACCENT
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = True
        tf_pill.margin_left = tf_pill.margin_top = tf_pill.margin_right = tf_pill.margin_bottom = 0
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = tag_text.upper()
        p_pill.font.size = Pt(9)
        p_pill.font.bold = True
        p_pill.font.color.rgb = C_CYAN_ACCENT if is_dark else C_INDIGO_ACCENT
        p_pill.alignment = PP_ALIGN.CENTER

        # Main Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.95))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_LIGHT if is_dark else C_TEXT_DARK
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(11.5)
            p2.font.color.rgb = C_TEXT_LIGHT_MUTED if is_dark else C_TEXT_MUTED
            p2.space_before = Pt(3)

    def add_footer(slide, slide_num, is_dark=False):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(8.5), Inches(0.35))
        tf = tx_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"IEEE Ignite Hackathon 2026  ·  Host: SKIT Jaipur  |  Team Capycoders  |  Problem ID: SKIT002  |  Slide {slide_num} of 10"
        p.font.size = Pt(9)
        p.font.color.rgb = C_TEXT_LIGHT_MUTED if is_dark else C_TEXT_MUTED
        p.alignment = PP_ALIGN.LEFT

        # Add GitHub link on right
        tx_box2 = slide.shapes.add_textbox(Inches(8.5), Inches(7.05), Inches(4.033), Inches(0.35))
        tf2 = tx_box2.text_frame
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = "github.com/Sunil-saini1818/MuleWatch"
        p2.font.size = Pt(9)
        p2.font.bold = True
        p2.font.color.rgb = C_CYAN_ACCENT if is_dark else C_INDIGO_ACCENT
        p2.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, title, items, badge="", accent_color=C_INDIGO_ACCENT, is_dark=False):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_NAVY_CARD if is_dark else C_CARD_BG
        card.line.color.rgb = C_NAVY_BORDER if is_dark else C_CARD_BORDER
        card.line.width = Pt(1)

        # Accent Top Stripe
        stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = accent_color
        stripe.line.fill.background()

        tx_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.2), width - Inches(0.44), height - Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        if badge:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge.upper()
            p_badge.font.size = Pt(9)
            p_badge.font.bold = True
            p_badge.font.color.rgb = accent_color
            p_badge.space_after = Pt(2)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_LIGHT if is_dark else C_TEXT_DARK
        p_title.space_after = Pt(7)

        for item in items:
            p_item = tf.add_paragraph()
            if isinstance(item, tuple):
                h, b = item
                run_h = p_item.add_run()
                run_h.text = f"• {h}: "
                run_h.font.bold = True
                run_h.font.size = Pt(10.5)
                run_h.font.color.rgb = C_TEXT_LIGHT if is_dark else C_TEXT_DARK
                run_b = p_item.add_run()
                run_b.text = b
                run_b.font.size = Pt(10)
                run_b.font.color.rgb = C_TEXT_LIGHT_MUTED if is_dark else C_TEXT_MUTED
            else:
                p_item.text = f"• {item}"
                p_item.font.size = Pt(10.5)
                p_item.font.color.rgb = C_TEXT_LIGHT_MUTED if is_dark else C_TEXT_MUTED
            p_item.space_after = Pt(5)

    def add_image_card(slide, left, top, width, height, img_path, caption=""):
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        frame.fill.solid()
        frame.fill.fore_color.rgb = C_CARD_BG
        frame.line.color.rgb = C_CARD_BORDER
        frame.line.width = Pt(1)

        pad = Inches(0.12)
        caption_h = Inches(0.35)
        img_avail_h = height - caption_h - (pad * 2)
        img_avail_w = width - (pad * 2)

        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, left + pad, top + pad, img_avail_w, img_avail_h)
        
        if caption:
            tx_box = slide.shapes.add_textbox(left + pad, top + height - caption_h - Inches(0.04), img_avail_w, caption_h)
            tf = tx_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = f"📷 {caption}"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = C_INDIGO_ACCENT
            p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, C_NAVY_DARK)

    # Top Event Banner Pill
    pill1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.7), Inches(5.8), Inches(0.4))
    pill1.fill.solid()
    pill1.fill.fore_color.rgb = C_NAVY_CARD
    pill1.line.color.rgb = C_INDIGO_LIGHT
    pill1.line.width = Pt(1)
    tf_p1 = pill1.text_frame
    tf_p1.margin_left = tf_p1.margin_top = tf_p1.margin_right = tf_p1.margin_bottom = 0
    p_p1 = tf_p1.paragraphs[0]
    p_p1.text = "⚡ IEEE IGNITE HACKATHON 2026  ·  HOST: SKIT JAIPUR"
    p_p1.font.size = Pt(10)
    p_p1.font.bold = True
    p_p1.font.color.rgb = C_CYAN_ACCENT
    p_p1.alignment = PP_ALIGN.CENTER

    # Main Project Title
    tx_title = slide1.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(1.8))
    tf_t = tx_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_main = tf_t.paragraphs[0]
    p_main.text = "MuleWatch"
    p_main.font.size = Pt(44)
    p_main.font.bold = True
    p_main.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf_t.add_paragraph()
    p_sub.text = "AI-Powered Detection of Suspicious Financial Advertisements & Cyber Cell Dispatch System"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = C_CYAN_ACCENT
    p_sub.space_before = Pt(6)

    p_desc = tf_t.add_paragraph()
    p_desc.text = "Autonomous threat intelligence harvesting, multimodal OCR/QR forensics, localized district risk mapping, and tamper-proof police evidence generation."
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = C_TEXT_LIGHT_MUTED
    p_desc.space_before = Pt(8)

    # 3 Metadata Cards
    m_top = Inches(3.6)
    m_w = Inches(3.7)
    m_h = Inches(1.9)

    card_meta1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), m_top, m_w, m_h)
    card_meta1.fill.solid()
    card_meta1.fill.fore_color.rgb = C_NAVY_CARD
    card_meta1.line.color.rgb = C_NAVY_BORDER
    tf_m1 = card_meta1.text_frame
    tf_m1.margin_left = tf_m1.margin_top = tf_m1.margin_right = tf_m1.margin_bottom = Inches(0.2)
    p_m1_tag = tf_m1.paragraphs[0]
    p_m1_tag.text = "HACKATHON CHALLENGE"
    p_m1_tag.font.size = Pt(9)
    p_m1_tag.font.bold = True
    p_m1_tag.font.color.rgb = C_CYAN_ACCENT
    p_m1_val = tf_m1.add_paragraph()
    p_m1_val.text = "Problem ID: SKIT002"
    p_m1_val.font.size = Pt(16)
    p_m1_val.font.bold = True
    p_m1_val.font.color.rgb = RGBColor(255, 255, 255)
    p_m1_sub = tf_m1.add_paragraph()
    p_m1_sub.text = "Theme: AI Detection of Suspicious Financial Ads & Mule Syndicates"
    p_m1_sub.font.size = Pt(10.5)
    p_m1_sub.font.color.rgb = C_TEXT_LIGHT_MUTED
    p_m1_sub.space_before = Pt(4)

    card_meta2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), m_top, m_w, m_h)
    card_meta2.fill.solid()
    card_meta2.fill.fore_color.rgb = C_NAVY_CARD
    card_meta2.line.color.rgb = C_NAVY_BORDER
    tf_m2 = card_meta2.text_frame
    tf_m2.margin_left = tf_m2.margin_top = tf_m2.margin_right = tf_m2.margin_bottom = Inches(0.2)
    p_m2_tag = tf_m2.paragraphs[0]
    p_m2_tag.text = "DEVELOPMENT TEAM"
    p_m2_tag.font.size = Pt(9)
    p_m2_tag.font.bold = True
    p_m2_tag.font.color.rgb = C_INDIGO_LIGHT
    p_m2_val = tf_m2.add_paragraph()
    p_m2_val.text = "Team: capycoders"
    p_m2_val.font.size = Pt(16)
    p_m2_val.font.bold = True
    p_m2_val.font.color.rgb = RGBColor(255, 255, 255)
    p_m2_sub = tf_m2.add_paragraph()
    p_m2_sub.text = "Focus: Full-Stack AI, Computer Vision, Digital Forensics & Cyber Defense"
    p_m2_sub.font.size = Pt(10.5)
    p_m2_sub.font.color.rgb = C_TEXT_LIGHT_MUTED
    p_m2_sub.space_before = Pt(4)

    card_meta3 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), m_top, m_w, m_h)
    card_meta3.fill.solid()
    card_meta3.fill.fore_color.rgb = C_NAVY_CARD
    card_meta3.line.color.rgb = C_NAVY_BORDER
    tf_m3 = card_meta3.text_frame
    tf_m3.margin_left = tf_m3.margin_top = tf_m3.margin_right = tf_m3.margin_bottom = Inches(0.2)
    p_m3_tag = tf_m3.paragraphs[0]
    p_m3_tag.text = "HOST INSTITUTION"
    p_m3_tag.font.size = Pt(9)
    p_m3_tag.font.bold = True
    p_m3_tag.font.color.rgb = C_GREEN_ACCENT
    p_m3_val = tf_m3.add_paragraph()
    p_m3_val.text = "SKIT Jaipur"
    p_m3_val.font.size = Pt(16)
    p_m3_val.font.bold = True
    p_m3_val.font.color.rgb = RGBColor(255, 255, 255)
    p_m3_sub = tf_m3.add_paragraph()
    p_m3_sub.text = "Swami Keshvanand Institute of Technology, Management & Gramothan"
    p_m3_sub.font.size = Pt(10.5)
    p_m3_sub.font.color.rgb = C_TEXT_LIGHT_MUTED
    p_m3_sub.space_before = Pt(4)

    # GitHub Banner at Bottom
    git_banner = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.75))
    git_banner.fill.solid()
    git_banner.fill.fore_color.rgb = C_NAVY_CARD
    git_banner.line.color.rgb = C_INDIGO_ACCENT
    git_banner.line.width = Pt(1)
    tf_git = git_banner.text_frame
    tf_git.margin_left = Inches(0.25)
    tf_git.margin_top = Inches(0.12)
    p_git = tf_git.paragraphs[0]
    run_git_label = p_git.add_run()
    run_git_label.text = "PUBLIC GITHUB REPOSITORY: "
    run_git_label.font.bold = True
    run_git_label.font.size = Pt(11)
    run_git_label.font.color.rgb = C_CYAN_ACCENT
    run_git_url = p_git.add_run()
    run_git_url.text = "https://github.com/Sunil-saini1818/MuleWatch  "
    run_git_url.font.size = Pt(11)
    run_git_url.font.bold = True
    run_git_url.font.color.rgb = RGBColor(255, 255, 255)
    run_git_badge = p_git.add_run()
    run_git_badge.text = "[Production-Ready Fullstack Codebase · Complete Commit History]"
    run_git_badge.font.size = Pt(10)
    run_git_badge.font.color.rgb = C_GREEN_ACCENT

    add_footer(slide1, 1, is_dark=True)

    # =========================================================================
    # SLIDE 2: Problem Statement & Background
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, C_BG_LIGHT)
    add_header(slide2, "PROBLEM ANALYSIS · THE THREAT LANDSCAPE",
               "The Silent Epidemic of Digital Mule Account Syndicates",
               "How organized financial cybercriminals exploit open communication channels and outpace manual policing")

    create_card(slide2, Inches(0.8), Inches(1.85), Inches(5.7), Inches(3.6),
                "The Underground Syndicate Model",
                [
                    ("Proliferation of Social Adverts", "Fraud rings openly advertise on Telegram, WhatsApp, Facebook, and forums offering cash for 'unused' bank accounts and UPI IDs."),
                    ("Commercialized Money Laundering", "Accounts are leased at 5% to 15% commission or purchased outright to layer proceeds from investment, gaming, and loan scams."),
                    ("Ultra-Rapid Churning & Evasion", "Syndicates discard channels, burn SIMs, and cycle UPI QR codes within hours of laundering dirty funds."),
                    ("Vulnerable Demographics Target", "Students and unemployed youths are enticed by 'easy work-from-home' earnings, unknowingly becoming criminal accomplices.")
                ],
                badge="Emerging Threat Vector",
                accent_color=C_RED_ACCENT)

    create_card(slide2, Inches(6.8), Inches(1.85), Inches(5.7), Inches(3.6),
                "Existing Law Enforcement Bottlenecks",
                [
                    ("Manual & Fragmented Monitoring", "Police personnel manually monitor social media—a reactive, unsustainable strategy facing millions of posts daily."),
                    ("Multimodal Obfuscation", "Fraudsters conceal UPI handles and bank details inside image banners and QR codes, completely evading keyword text filters."),
                    ("48–72 Hour Critical Triage Lag", "By the time a complaint is registered, mule accounts have dispersed funds across multiple layered banking layers."),
                    ("Lack of Court-Admissible Evidence", "Raw phone screenshots lack cryptographic timestamps, origin IP metadata, and tamper-proof SHA-256 digital hashes required for FIRs.")
                ],
                badge="Operational Deficiency",
                accent_color=C_INDIGO_ACCENT)

    # Bottom Stat Ribbon
    ribbon = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.65), Inches(11.733), Inches(1.15))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_CARD_BG
    ribbon.line.color.rgb = C_CARD_BORDER
    tf_rib = ribbon.text_frame
    tf_rib.margin_top = Inches(0.12)
    tf_rib.margin_left = Inches(0.2)
    p_rib_t = tf_rib.paragraphs[0]
    p_rib_t.text = "THE MULE ACCOUNT PIPELINE: HOW ORCHESTRATED CRIME DRAINS CITIZEN SAVINGS"
    p_rib_t.font.size = Pt(9.5)
    p_rib_t.font.bold = True
    p_rib_t.font.color.rgb = C_RED_ACCENT
    p_rib_b = tf_rib.add_paragraph()
    p_rib_b.text = "1. Social Media Ad Recruiting Mules  ➔  2. KYC & UPI Transfer via Telegram  ➔  3. Scammed Funds Injected  ➔  4. Instant ATM / Crypto Cashout (Zero Trace)"
    p_rib_b.font.size = Pt(11)
    p_rib_b.font.bold = True
    p_rib_b.font.color.rgb = C_TEXT_DARK
    p_rib_b.space_before = Pt(4)

    add_footer(slide2, 2)

    # =========================================================================
    # SLIDE 3: Solution Overview (MuleWatch)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, C_BG_LIGHT)
    add_header(slide3, "SYSTEM SOLUTION · EXECUTIVE OVERVIEW",
               "MuleWatch: Autonomous Threat Intelligence to Cyber Interdiction",
               "An integrated fullstack platform bridging continuous threat crawling, multimodal AI, and rapid cyber cell dispatch")

    create_card(slide3, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Key Capabilities & Architecture",
                [
                    ("Autonomous Multichannel Harvesting", "Continuous background ingestion workers monitor Telegram underground groups, discussion forums (Charcha), and public social streams."),
                    ("Multimodal Entity Extraction", "Combines high-accuracy Tesseract OCR with OpenCV QR decoding to extract masked bank numbers, UPI handles, and payment links."),
                    ("Dynamic Heuristic Risk Scoring", "NLP threat categorization instantly scores posts (0-100%) against urgency triggers, commission bait, and illicit financial keywords."),
                    ("Mule Syndicate Entity Linking", "Resolves disparate posts into unified actor profiles, linking repeated UPI handles, phone numbers, and associated telegram aliases."),
                    ("Cyber Cell Dispatch & I4C Webhooks", "1-click generation of court-admissible PDF dossiers with SHA-256 tamper-proof hashes and direct dispatch to bank nodal officers.")
                ],
                badge="End-to-End Cyber Defense",
                accent_color=C_INDIGO_ACCENT)

    add_image_card(slide3, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_DASHBOARD, "Figure 1: MuleWatch Real-Time Command Center & Threat Stream")

    add_footer(slide3, 3)

    # =========================================================================
    # SLIDE 4: Core Architecture & Tech Stack
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, C_BG_LIGHT)
    add_header(slide4, "SYSTEM ARCHITECTURE · TECHNOLOGY STACK",
               "Production-Grade, Modular & Resilient Fullstack Stack",
               "Engineered for high-concurrency ingestion, asynchronous processing, and zero-compromise digital forensics")

    create_card(slide4, Inches(0.8), Inches(1.85), Inches(3.7), Inches(4.9),
                "Frontend Experience Layer",
                [
                    ("React 18 & Vite", "Lightning-fast reactive UI with hot module replacement and optimized production bundle."),
                    ("Tailwind CSS & Lucide Icons", "Ergonomic, modern cyber operations center aesthetic with high-contrast data visualization."),
                    ("Gapless TopoJSON Vector Map", "Custom-rendered SVG district map of Rajasthan with zero-gap borders and hover state binding."),
                    ("Dynamic Polling Telemetry", "Self-healing telemetry hooks polling live crawler status, alert counts, and actor networks.")
                ],
                badge="Client Presentation",
                accent_color=C_INDIGO_ACCENT)

    create_card(slide4, Inches(4.8), Inches(1.85), Inches(3.7), Inches(4.9),
                "Backend & Intelligence Core",
                [
                    ("FastAPI & Python 3.12", "Asynchronous ASGI framework delivering ultra-low latency REST APIs and strict Pydantic models."),
                    ("Computer Vision & OCR", "Tesseract OCR with adaptive image binarization; PyZBar for instant UPI QR payload decoding."),
                    ("Telethon MTProto Client", "Native Telegram API integration capable of real-time message stream harvesting."),
                    ("SQLite WAL Database", "High-performance Write-Ahead Logging mode providing ACID transactions without heavy database overhead.")
                ],
                badge="Microservice Engine",
                accent_color=C_NAVY_DARK)

    create_card(slide4, Inches(8.8), Inches(1.85), Inches(3.733), Inches(4.9),
                "Forensics & Police Interdiction",
                [
                    ("SHA-256 Evidence Seals", "Cryptographic fingerprinting of raw ad screenshots and metadata to ensure Section 65B legal validity."),
                    ("Automated PDF Dossiers", "ReportLab PDF engine compiling investigator-ready case reports complete with IP and user-agent data."),
                    ("Reversible State Machine", "Enforcement safety guardrail allowing investigators to confirm fraud or revert false positives seamlessly."),
                    ("Unified Launcher (`./run.sh`)", "One-click shell deployment bootstrapping backend, frontend, and simulated mock services concurrently.")
                ],
                badge="Forensic Integrity",
                accent_color=C_GREEN_ACCENT)

    add_footer(slide4, 4)

    # =========================================================================
    # SLIDE 5: Key Feature 1 - Multi-Platform Crawling & Ingestion
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, C_BG_LIGHT)
    add_header(slide5, "CORE CAPABILITY 1 · INTELLIGENCE HARVESTING",
               "Autonomous Multi-Source Ingestion & Crawler Telemetry",
               "Proactively scanning underground Telegram channels, public discussion boards, and social threads")

    create_card(slide5, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Automated Ingestion Infrastructure",
                [
                    ("Telethon MTProto Telegram Ingestion", "Subscribes directly to suspect Telegram channels, capturing incoming messages, forwarded posts, and embedded image assets in real time."),
                    ("Charcha Discussion Forum Scraper", "Dedicated resilient crawler with dynamic pagination, thread tracking, and deduplication for local underground discussion boards."),
                    ("SimHash & LSH Deduplication", "Locality-Sensitive Hashing identifies duplicate syndicated ads reposted across dozens of channels, consolidating threat records."),
                    ("Live Crawler Health Telemetry", "Real-time heartbeat monitoring tracks active endpoints, fetch latencies, and total harvested records with automatic retry policies."),
                    ("Extensible Ingestion API", "Standardized webhook receiver accepts external threat feeds from partner cybersecurity teams and public reporting portals.")
                ],
                badge="Data Ingestion Pipeline",
                accent_color=C_INDIGO_ACCENT)

    add_image_card(slide5, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_LINKED_PLATFORMS, "Figure 2: Multi-Platform Ingestion Telemetry & Health Monitoring")

    add_footer(slide5, 5)

    # =========================================================================
    # SLIDE 6: Key Feature 2 - Advanced Multimodal OCR & QR Analysis
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, C_BG_LIGHT)
    add_header(slide6, "CORE CAPABILITY 2 · MULTIMODAL FORENSICS",
               "Multimodal OCR, QR Code Decoding & Entity Extraction",
               "Exposing financial identifiers deliberately disguised within images, posters, and payment barcodes")

    create_card(slide6, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Deep Visual Extraction Engine",
                [
                    ("Adaptive Image Pre-Processing", "Applies contrast stretching, adaptive thresholding, and noise reduction to overcome obfuscated fonts and low-resolution screenshots."),
                    ("Tesseract OCR Entity Parsing", "Extracts masked bank names, IFSC codes, phone numbers, and commission percentages from graphic ad banners."),
                    ("Instant UPI & Crypto QR Decoding", "OpenCV and PyZBar decoder extracts raw `upi://pay?pa=...` merchant strings and crypto wallet addresses directly from payment graphics."),
                    ("NLP Risk Scoring Heuristics", "Evaluates linguistic urgency indicators ('Immediate Need', 'Rent Bank Account', '10% Daily Commission') into weighted threat scores."),
                    ("Interactive 'Check a Post' Hub", "Allows investigators to test suspicious URLs, paste raw text, or upload image screenshots for instantaneous risk analysis.")
                ],
                badge="Multimodal Intelligence",
                accent_color=C_INDIGO_ACCENT)

    add_image_card(slide6, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_CHECK_POST, "Figure 3: Interactive Multimodal Analyzer with QR & Entity Extraction")

    add_footer(slide6, 6)

    # =========================================================================
    # SLIDE 7: Regional Intelligence - Rajasthan District Fraud Map
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, C_BG_LIGHT)
    add_header(slide7, "CORE CAPABILITY 3 · GEOSPATIAL INTELLIGENCE",
               "Localized Rajasthan District-Level Fraud Analytics",
               "Interactive TopoJSON visualization mapping mule account syndicates across Rajasthan administrative districts")

    create_card(slide7, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Hyper-Local Regional Telemetry",
                [
                    ("High-Fidelity Vector TopoJSON", "Custom-crafted SVG district boundaries with clean, gapless polygon borders, eliminating geographic rendering distortion."),
                    ("Real-World Hotspot Correlation", "Reflects real ground intelligence across known financial cybercrime hotspots including Jaipur, Bharatpur, Alwar, and Mewat border zones."),
                    ("Dynamic Hover & Click Data Binding", "Interactive UI binds to live district metrics, displaying active cases, suspected mule accounts, flagged volume, and resolution rates."),
                    ("Domain-Specific Analytics Table", "Replaces generic loan metrics with actionable cyber metrics: 'Confirmed Mule Cases', 'Suspected Bank Accounts', and 'Lien Requests'."),
                    ("Tactical Law Enforcement Deployment", "Enables State Cyber Police Headquarters to allocate investigative resources directly to emerging fraud epicenters.")
                ],
                badge="Geospatial Threat Mapping",
                accent_color=C_INDIGO_ACCENT)

    add_image_card(slide7, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_RAJASTHAN_MAP, "Figure 4: Rajasthan District-Level Fraud TopoJSON Intelligence Map")

    add_footer(slide7, 7)

    # =========================================================================
    # SLIDE 8: Enforcement Workflow - Cyber Cell Dispatch
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, C_BG_LIGHT)
    add_header(slide8, "CORE CAPABILITY 4 · LAW ENFORCEMENT INTERDICTION",
               "Cyber Cell Dispatch & Tamper-Proof Evidence Dossier",
               "Bridging the critical gap between automated intelligence and judicial action with 1-click police exports")

    create_card(slide8, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Operational Enforcement Pipeline",
                [
                    ("Dedicated Nodal Dispatch Queue", "Confirmed fraud cases move automatically from general investigation into a dedicated 'Cyber Cell Dispatch' operational queue."),
                    ("Automated Evidence Dossier (PDF)", "Generates standardized judicial evidence dossiers featuring post URLs, origin IP addresses, actor handles, and device user-agents."),
                    ("Cryptographic SHA-256 Proof", "Every generated report seals raw evidence with SHA-256 checksums to satisfy Section 65B Indian Evidence Act standards."),
                    ("1-Click Simulated Dispatches", "Direct webhook dispatches to the National Cyber Crime Reporting Portal (I4C) and automated bank account lien freeze notices."),
                    ("Reversal Guardrails ('Revert to Draft')", "Built-in safety state machine allows officers to instantly reverse accidental confirmations with zero audit trail corruption.")
                ],
                badge="Judicial & Banking Interdiction",
                accent_color=C_GREEN_ACCENT)

    add_image_card(slide8, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_PDF_REPORT, "Figure 5: Forensic Evidence Report with SHA-256 Cryptographic Hash")

    add_footer(slide8, 8)

    # =========================================================================
    # SLIDE 9: Impact, Network Intelligence & Forensic Audit
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, C_BG_LIGHT)
    add_header(slide9, "IMPACT & FORENSIC SYNERGY · OPERATIONAL METRICS",
               "Measurable Impact & Syndicate Network Linkage",
               "Uncovering complex multi-actor mule syndicates through automated entity relationship graphs")

    create_card(slide9, Inches(0.8), Inches(1.85), Inches(5.3), Inches(4.9),
                "Quantifiable Operational Advantage",
                [
                    (">85% Reduction in Triage Time", "Reduces incident verification timelines from 48-72 hours of manual searching to under 3 minutes per flagged advertisement."),
                    ("Mule Syndicate Linkage Graph", "Force-directed visual graph uncovers shared UPI handles, phone numbers, and Telegram channels linking disparate mule accounts."),
                    ("Digital Chain of Custody", "Immutable digital audit trail records timestamps, examiner actions, and state transitions for courtroom presentation."),
                    ("Inter-Agency Collaboration", "Standardized JSON/PDF export formats allow seamless intelligence exchange between State Cyber Cells and Nodal Bank Officers."),
                    ("Proactive Threat Neutralization", "Freezes mule bank accounts before dirty scam funds can be layered or cashed out at ATMs.")
                ],
                badge="High-Impact Outcomes",
                accent_color=C_INDIGO_ACCENT)

    add_image_card(slide9, Inches(6.3), Inches(1.85), Inches(6.233), Inches(4.9),
                   IMG_SELLER_NETWORK, "Figure 6: Mule Syndicate Linkage Graph & Financial Entity Clustering")

    add_footer(slide9, 9)

    # =========================================================================
    # SLIDE 10: Conclusion, Future Roadmap & Q&A (Dark Theme)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, C_NAVY_DARK)
    add_header(slide10, "MISSION SUMMARY · FUTURE HORIZONS",
               "Securing the Digital Financial Ecosystem",
               "MuleWatch delivers an operational, end-to-end defense system against organized financial cybercrime",
               is_dark=True)

    create_card(slide10, Inches(0.8), Inches(1.85), Inches(3.7), Inches(3.7),
                "Delivered Milestones",
                [
                    ("Production Fullstack Engine", "FastAPI + React 18 with persistent SQLite WAL database and live telemetry."),
                    ("Multimodal AI Pipeline", "Automated OCR extraction, QR decoding, and NLP heuristic threat scoring."),
                    ("Rajasthan District Map", "Gapless TopoJSON SVG vector map bound to localized district metrics."),
                    ("Cyber Cell Dispatch Queue", "1-click I4C dispatch, bank lien notices, and SHA-256 PDF dossiers.")
                ],
                badge="Hackathon Execution",
                accent_color=C_GREEN_ACCENT,
                is_dark=True)

    create_card(slide10, Inches(4.8), Inches(1.85), Inches(3.7), Inches(3.7),
                "Immediate Roadmap",
                [
                    ("WhatsApp & Dark Web Crawlers", "Expanding autonomous crawlers to WhatsApp Business channels and Tor forums."),
                    ("Direct NPCI / RBI API Hooks", "Real-time automated account freezing via official banking nodal gateways."),
                    ("Indic LLM Dialect Fine-Tuning", "Specialized NLP models for regional Indian slang, Romanized Hindi, and evasive code-words."),
                    ("Automated Sec 65B Certificates", "Automated digital signing and hash stamping for formal court filings.")
                ],
                badge="Production Scale",
                accent_color=C_CYAN_ACCENT,
                is_dark=True)

    create_card(slide10, Inches(8.8), Inches(1.85), Inches(3.733), Inches(3.7),
                "Team & Repository",
                [
                    ("Team Capycoders", "Developers dedicated to building high-impact cyber defense and financial integrity tooling."),
                    ("Host Institution", "SKIT Jaipur (Swami Keshvanand Institute of Technology)"),
                    ("Event Context", "IEEE Ignite Hackathon 2026 · Problem ID: SKIT002"),
                    ("GitHub Repository", "github.com/Sunil-saini1818/MuleWatch (Public, active commit history)")
                ],
                badge="Open Source & Ready",
                accent_color=C_INDIGO_LIGHT,
                is_dark=True)

    # Q&A Banner at bottom
    qa_banner = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.05))
    qa_banner.fill.solid()
    qa_banner.fill.fore_color.rgb = C_NAVY_CARD
    qa_banner.line.color.rgb = C_CYAN_ACCENT
    qa_banner.line.width = Pt(1.5)
    tf_qa = qa_banner.text_frame
    tf_qa.margin_top = Inches(0.15)
    tf_qa.margin_left = Inches(0.3)
    p_qa1 = tf_qa.paragraphs[0]
    p_qa1.text = "THANK YOU  ·  READY FOR LIVE DEMONSTRATION & JUDGES' Q&A"
    p_qa1.font.size = Pt(14)
    p_qa1.font.bold = True
    p_qa1.font.color.rgb = RGBColor(255, 255, 255)
    p_qa2 = tf_qa.add_paragraph()
    p_qa2.text = "Explore the live system, test suspicious advertisements, and verify the end-to-end evidence generation pipeline."
    p_qa2.font.size = Pt(11)
    p_qa2.font.color.rgb = C_TEXT_LIGHT_MUTED
    p_qa2.space_before = Pt(3)

    add_footer(slide10, 10, is_dark=True)

    # Save presentation
    output_filename = "MuleWatch_Presentation_IEEE_Ignite_2026.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'")

    # Also save to artifact directory
    artifact_dir = "/Users/mahak/.gemini/antigravity/brain/63ff7a6c-88ed-420f-957e-9fb7f22791b3"
    artifact_path = os.path.join(artifact_dir, output_filename)
    prs.save(artifact_path)
    print(f"Presentation also saved to artifact directory: '{artifact_path}'")

if __name__ == "__main__":
    build_presentation()
